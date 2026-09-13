"""
Comprehensive test suite for 56 File Converter engines.
"""

import os
import shutil
import unittest
from PIL import Image
import docx
import openpyxl

from converters.registry import convert_file, get_supported_targets, normalize_ext

TEST_DIR = os.path.join(os.path.dirname(__file__), "scratch_test_files")
OUT_DIR = os.path.join(os.path.dirname(__file__), "scratch_test_out")


class TestConverters(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        os.makedirs(TEST_DIR, exist_ok=True)
        os.makedirs(OUT_DIR, exist_ok=True)

        # 1. Create a test DOCX
        cls.docx_path = os.path.join(TEST_DIR, "sample.docx")
        doc = docx.Document()
        doc.add_heading("56 File Converter Test", level=1)
        doc.add_paragraph("This is a test paragraph for verifying document conversions.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Data A"
        table.cell(1, 1).text = "Data B"
        doc.save(cls.docx_path)

        # 2. Create a test Image (PNG)
        cls.png_path = os.path.join(TEST_DIR, "sample.png")
        img = Image.new("RGBA", (120, 120), (79, 70, 229, 255))
        img.save(cls.png_path)

        # 3. Create a test TXT
        cls.txt_path = os.path.join(TEST_DIR, "sample.txt")
        with open(cls.txt_path, "w", encoding="utf-8") as f:
            f.write("Hello from 56 File Converter!\nTesting text conversion pipeline.")

        # 4. Create a test MD
        cls.md_path = os.path.join(TEST_DIR, "sample.md")
        with open(cls.md_path, "w", encoding="utf-8") as f:
            f.write("# Title\n\n- Feature 1\n- Feature 2\n\n```python\nprint('hello')\n```")

        # 5. Create a test CSV
        cls.csv_path = os.path.join(TEST_DIR, "sample.csv")
        with open(cls.csv_path, "w", encoding="utf-8") as f:
            f.write("Name,Role,Country\nAlice,Developer,USA\nBob,Designer,UK\n")

        # 6. Create a test XLSX
        cls.xlsx_path = os.path.join(TEST_DIR, "sample.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Product", "Price", "Stock"])
        ws.append(["Widget", 25, 100])
        ws.append(["Gadget", 45, 50])
        wb.save(cls.xlsx_path)

        # 7. Create a test JSON
        cls.json_path = os.path.join(TEST_DIR, "sample.json")
        with open(cls.json_path, "w", encoding="utf-8") as f:
            f.write('[{"id": 1, "name": "Item 1"}, {"id": 2, "name": "Item 2"}]')

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TEST_DIR, ignore_errors=True)
        shutil.rmtree(OUT_DIR, ignore_errors=True)

    def test_registry_supported_targets(self):
        targets = get_supported_targets("docx")
        exts = [t["extension"] for t in targets]
        self.assertIn("pdf", exts)
        self.assertIn("txt", exts)

        targets_jpg = get_supported_targets("jpg")
        exts_jpg = [t["extension"] for t in targets_jpg]
        self.assertIn("png", exts_jpg)
        self.assertIn("webp", exts_jpg)
        self.assertIn("pdf", exts_jpg)

    def test_docx_to_pdf_and_back(self):
        # DOCX -> PDF
        pdf_out = convert_file(self.docx_path, "docx", "pdf", OUT_DIR)
        self.assertTrue(os.path.exists(pdf_out))
        self.assertTrue(os.path.getsize(pdf_out) > 0)

        # PDF -> DOCX
        docx_out = convert_file(pdf_out, "pdf", "docx", OUT_DIR)
        self.assertTrue(os.path.exists(docx_out))
        self.assertTrue(os.path.getsize(docx_out) > 0)

        # PDF -> PNG
        png_out = convert_file(pdf_out, "pdf", "png", OUT_DIR)
        self.assertTrue(os.path.exists(png_out))
        self.assertTrue(os.path.getsize(png_out) > 0)

        # PDF -> TXT
        txt_out = convert_file(pdf_out, "pdf", "txt", OUT_DIR)
        self.assertTrue(os.path.exists(txt_out))
        self.assertTrue(os.path.getsize(txt_out) > 0)

    def test_image_conversions(self):
        # PNG -> JPG
        jpg_out = convert_file(self.png_path, "png", "jpg", OUT_DIR)
        self.assertTrue(os.path.exists(jpg_out))

        # JPG -> WEBP
        webp_out = convert_file(jpg_out, "jpg", "webp", OUT_DIR)
        self.assertTrue(os.path.exists(webp_out))

        # PNG -> PDF
        img_pdf_out = convert_file(self.png_path, "png", "pdf", OUT_DIR)
        self.assertTrue(os.path.exists(img_pdf_out))

        # PNG -> ICO
        ico_out = convert_file(self.png_path, "png", "ico", OUT_DIR)
        self.assertTrue(os.path.exists(ico_out))

    def test_data_conversions(self):
        # CSV -> XLSX
        xlsx_out = convert_file(self.csv_path, "csv", "xlsx", OUT_DIR)
        self.assertTrue(os.path.exists(xlsx_out))

        # CSV -> JSON
        json_out = convert_file(self.csv_path, "csv", "json", OUT_DIR)
        self.assertTrue(os.path.exists(json_out))

        # XLSX -> CSV
        csv_out = convert_file(self.xlsx_path, "xlsx", "csv", OUT_DIR)
        self.assertTrue(os.path.exists(csv_out))

        # JSON -> CSV
        json_csv_out = convert_file(self.json_path, "json", "csv", OUT_DIR)
        self.assertTrue(os.path.exists(json_csv_out))

    def test_markdown_and_text(self):
        # MD -> HTML
        html_out = convert_file(self.md_path, "md", "html", OUT_DIR)
        self.assertTrue(os.path.exists(html_out))

        # MD -> PDF
        pdf_out = convert_file(self.md_path, "md", "pdf", OUT_DIR)
        self.assertTrue(os.path.exists(pdf_out))

        # TXT -> PDF
        txt_pdf_out = convert_file(self.txt_path, "txt", "pdf", OUT_DIR)
        self.assertTrue(os.path.exists(txt_pdf_out))

    def test_empty_and_unusual_docx(self):
        """Verify converter gracefully handles empty or unusual DOCX files without crashing."""
        empty_docx = os.path.join(TEST_DIR, "empty.docx")
        doc = docx.Document()
        doc.save(empty_docx)

        # Empty docx to PDF
        pdf_out = convert_file(empty_docx, "docx", "pdf", OUT_DIR)
        self.assertTrue(os.path.exists(pdf_out))
        self.assertTrue(os.path.getsize(pdf_out) > 0)

        # Empty docx to TXT
        txt_out = convert_file(empty_docx, "docx", "txt", OUT_DIR)
        self.assertTrue(os.path.exists(txt_out))

        # Empty docx to HTML
        html_out = convert_file(empty_docx, "docx", "html", OUT_DIR)
        self.assertTrue(os.path.exists(html_out))


if __name__ == "__main__":
    unittest.main()
