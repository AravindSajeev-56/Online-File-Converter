"""
Document Converter Module for 56 File Converter
Handles DOCX, PDF, TXT, MD, and HTML conversions.
"""

import os
import re
import zipfile
import html
from xml.sax.saxutils import escape
from typing import List

# Third-party libraries
import pymupdf
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
import docx
import markdown


def _sanitize_text_for_reportlab(text: str) -> str:
    """Escape XML characters and clean text for ReportLab Paragraphs."""
    if not text:
        return ""
    # Strip non-printable/control characters that break ReportLab
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    escaped = html.escape(text)
    return escaped


def _build_pdf_from_elements(story: list, output_path: str):
    """Build a standard PDF from Platypus flowable elements."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    doc.build(story)


def docx_to_pdf(input_path: str, output_path: str) -> str:
    """Convert DOCX file to PDF using python-docx and reportlab."""
    source_doc = docx.Document(input_path)
    styles = getSampleStyleSheet()
    
    # Custom styles
    body_style = ParagraphStyle(
        'DocxBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor('#1f2937'),
        spaceAfter=8
    )
    h1_style = ParagraphStyle(
        'DocxH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#111827'),
        spaceBefore=12,
        spaceAfter=10
    )
    h2_style = ParagraphStyle(
        'DocxH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#1f2937'),
        spaceBefore=10,
        spaceAfter=6
    )
    h3_style = ParagraphStyle(
        'DocxH3',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#374151'),
        spaceBefore=8,
        spaceAfter=4
    )

    story = []
    
    # Process document body
    for element in source_doc.element.body:
        tag = element.tag.split('}')[-1]
        
        if tag == 'p':
            # It's a paragraph
            p = docx.text.paragraph.Paragraph(element, source_doc)
            raw_text = p.text.strip()
            if not raw_text:
                story.append(Spacer(1, 6))
                continue
            
            clean_text = _sanitize_text_for_reportlab(raw_text)
            p_style = p.style.name.lower() if p.style and p.style.name else ""
            
            if 'heading 1' in p_style or 'title' in p_style:
                story.append(Paragraph(clean_text, h1_style))
            elif 'heading 2' in p_style:
                story.append(Paragraph(clean_text, h2_style))
            elif 'heading 3' in p_style:
                story.append(Paragraph(clean_text, h3_style))
            else:
                story.append(Paragraph(clean_text, body_style))
                
        elif tag == 'tbl':
            # It's a table
            table = docx.table.Table(element, source_doc)
            table_data = []
            for row in table.rows:
                row_cells = []
                for cell in row.cells:
                    row_cells.append(Paragraph(_sanitize_text_for_reportlab(cell.text.strip()), body_style))
                table_data.append(row_cells)
                
            if table_data:
                col_count = max(len(row) for row in table_data) if table_data else 1
                for row in table_data:
                    while len(row) < col_count:
                        row.append(Paragraph("", body_style))
                available_width = letter[0] - 108
                col_width = available_width / max(col_count, 1)
                
                t = Table(table_data, colWidths=[col_width] * col_count)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f3f4f6')),
                    ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1f2937')),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#d1d5db')),
                ]))
                story.append(Spacer(1, 8))
                story.append(t)
                story.append(Spacer(1, 8))

    if not story:
        story.append(Paragraph("Empty document", body_style))

    try:
        _build_pdf_from_elements(story, output_path)
    except Exception:
        # Fallback text rendering
        fallback_story = []
        for p in source_doc.paragraphs:
            txt = _sanitize_text_for_reportlab(p.text.strip())
            if txt:
                fallback_story.append(Paragraph(txt, body_style))
                fallback_story.append(Spacer(1, 6))
        if not fallback_story:
            fallback_story.append(Paragraph("Empty document", body_style))
        _build_pdf_from_elements(fallback_story, output_path)

    return output_path


def pdf_to_docx(input_path: str, output_path: str) -> str:
    """Convert PDF file to DOCX using pdf2docx with PyMuPDF text fallback."""
    try:
        from pdf2docx import Converter
        cv = Converter(input_path)
        try:
            cv.convert(output_path)
        finally:
            cv.close()
            import gc
            gc.collect()
        if os.path.exists(output_path) and os.path.getsize(output_path) > 0:
            return output_path
    except Exception:
        pass

    # Fallback: extract pages text via PyMuPDF into docx
    doc = pymupdf.open(input_path)
    try:
        new_doc = docx.Document()
        for i, page in enumerate(doc):
            if i > 0:
                new_doc.add_page_break()
            txt = page.get_text()
            if txt.strip():
                for line in txt.splitlines():
                    if line.strip():
                        new_doc.add_paragraph(line)
        new_doc.save(output_path)
    finally:
        doc.close()
        import gc
        gc.collect()
    return output_path


def pdf_to_txt(input_path: str, output_path: str) -> str:
    """Extract text from PDF using pymupdf."""
    doc = pymupdf.open(input_path)
    text_content = []
    for page_num, page in enumerate(doc, 1):
        page_text = page.get_text()
        text_content.append(f"--- Page {page_num} ---\n" + page_text)
    doc.close()
    
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(text_content))
    return output_path


def pdf_to_images(input_path: str, output_dir: str, target_format: str = "png") -> str:
    """Convert PDF pages to image(s). If multiple pages, returns a ZIP file."""
    doc = pymupdf.open(input_path)
    page_count = len(doc)
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    
    if page_count == 1:
        out_file = os.path.join(output_dir, f"{base_name}.{target_format}")
        page = doc[0]
        pix = page.get_pixmap(dpi=200)
        pix.save(out_file)
        doc.close()
        return out_file
    else:
        image_files = []
        for i, page in enumerate(doc, 1):
            img_name = f"{base_name}_page_{i}.{target_format}"
            img_path = os.path.join(output_dir, img_name)
            pix = page.get_pixmap(dpi=200)
            pix.save(img_path)
            image_files.append((img_path, img_name))
        doc.close()
        
        # Package into zip
        zip_path = os.path.join(output_dir, f"{base_name}_pages.zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zipf:
            for file_path, arcname in image_files:
                zipf.write(file_path, arcname)
                try:
                    os.remove(file_path)
                except OSError:
                    pass
        return zip_path


def pdf_to_html(input_path: str, output_path: str) -> str:
    """Convert PDF to HTML layout using pymupdf."""
    doc = pymupdf.open(input_path)
    html_pages = []
    for i, page in enumerate(doc, 1):
        html_pages.append(f"<div class='pdf-page' id='page-{i}'>\n" + page.get_text("html") + "\n</div>")
    doc.close()
    
    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Converted PDF</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #f9fafb; padding: 20px; }}
  .pdf-page {{ background: white; margin: 20px auto; padding: 40px; max-width: 800px; box-shadow: 0 4px 6px rgba(0,0,0,0.1); border-radius: 8px; }}
</style>
</head>
<body>
{"<hr style='border:0;border-top:1px dashed #ccc;margin:30px 0;'>".join(html_pages)}
</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    return output_path


def docx_to_txt(input_path: str, output_path: str) -> str:
    """Extract plain text from DOCX."""
    doc = docx.Document(input_path)
    lines = []
    for p in doc.paragraphs:
        lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return output_path


def docx_to_html(input_path: str, output_path: str) -> str:
    """Convert DOCX to HTML document."""
    doc = docx.Document(input_path)
    body_html = []
    for p in doc.paragraphs:
        txt = html.escape(p.text)
        if not txt.strip():
            continue
        style = p.style.name.lower() if p.style and p.style.name else ""
        if 'heading 1' in style or 'title' in style:
            body_html.append(f"<h1>{txt}</h1>")
        elif 'heading 2' in style:
            body_html.append(f"<h2>{txt}</h2>")
        elif 'heading 3' in style:
            body_html.append(f"<h3>{txt}</h3>")
        else:
            body_html.append(f"<p>{txt}</p>")

    for table in doc.tables:
        rows_html = []
        for row in table.rows:
            cells = "".join(f"<td>{html.escape(c.text.strip())}</td>" for c in row.cells)
            rows_html.append(f"<tr>{cells}</tr>")
        body_html.append(f"<table border='1' cellpadding='6' cellspacing='0'>{''.join(rows_html)}</table>")

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Converted Document</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; max-width: 800px; margin: 40px auto; padding: 20px; line-height: 1.6; color: #111827; }}
  table {{ width: 100%; border-collapse: collapse; margin: 20px 0; }}
  th, td {{ border: 1px solid #d1d5db; padding: 8px 12px; }}
</style>
</head>
<body>
{"".join(body_html)}
</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    return output_path


def txt_to_pdf(input_path: str, output_path: str) -> str:
    """Convert plain text file to PDF."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    
    styles = getSampleStyleSheet()
    pre_style = ParagraphStyle(
        'CodeStyle',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1f2937')
    )
    
    story = []
    chunks = text.split("\n\n")
    for chunk in chunks:
        clean = _sanitize_text_for_reportlab(chunk)
        if clean.strip():
            story.append(Paragraph(clean.replace("\n", "<br/>"), pre_style))
            story.append(Spacer(1, 8))
            
    if not story:
        story.append(Paragraph("Empty text file", pre_style))
        
    _build_pdf_from_elements(story, output_path)
    return output_path


def txt_to_docx(input_path: str, output_path: str) -> str:
    """Convert text file to DOCX."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    doc = docx.Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    doc.save(output_path)
    return output_path


def txt_to_html(input_path: str, output_path: str) -> str:
    """Convert text file to HTML."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    escaped = html.escape(text)
    full_html = f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><title>Converted Text</title>
<style>body {{ font-family: monospace; padding: 30px; white-space: pre-wrap; line-height: 1.5; background: #fafafa; color: #222; }}</style>
</head>
<body>{escaped}</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    return output_path


def md_to_html(input_path: str, output_path: str) -> str:
    """Convert Markdown to styled HTML."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        md_text = f.read()
    html_content = markdown.markdown(md_text, extensions=['extra', 'codehilite', 'toc'])
    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Converted Markdown</title>
<style>
  body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; max-width: 820px; margin: 40px auto; padding: 24px; line-height: 1.7; color: #1e293b; background: #fff; }}
  pre {{ background: #f1f5f9; padding: 14px; border-radius: 6px; overflow-x: auto; }}
  code {{ font-family: monospace; background: #f1f5f9; padding: 2px 6px; border-radius: 4px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 20px 0; }}
  th, td {{ border: 1px solid #cbd5e1; padding: 8px 12px; text-align: left; }}
  th {{ background: #f8fafc; }}
  blockquote {{ border-left: 4px solid #6366f1; margin: 0; padding-left: 16px; color: #475569; }}
</style>
</head>
<body>
{html_content}
</body>
</html>"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_html)
    return output_path


def md_to_pdf(input_path: str, output_path: str) -> str:
    """Convert Markdown to PDF via reportlab."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        md_text = f.read()
    
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        'MDBody', parent=styles['Normal'], fontSize=10, leading=14, spaceAfter=8
    )
    h1_style = ParagraphStyle(
        'MDH1', parent=styles['Heading1'], fontSize=18, leading=22, spaceBefore=12, spaceAfter=8
    )
    h2_style = ParagraphStyle(
        'MDH2', parent=styles['Heading2'], fontSize=14, leading=18, spaceBefore=10, spaceAfter=6
    )
    code_style = ParagraphStyle(
        'MDCode', parent=styles['Normal'], fontName='Courier', fontSize=9, leading=12,
        textColor=colors.HexColor('#0f172a'), spaceAfter=6
    )
    
    story = []
    lines = md_text.splitlines()
    in_code_block = False
    code_lines = []
    
    for line in lines:
        if line.strip().startswith("```"):
            if in_code_block:
                in_code_block = False
                story.append(Preformatted("\n".join(code_lines), code_style))
                story.append(Spacer(1, 6))
                code_lines = []
            else:
                in_code_block = True
            continue
        
        if in_code_block:
            code_lines.append(line)
            continue
            
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 6))
            continue
            
        if stripped.startswith("# "):
            story.append(Paragraph(_sanitize_text_for_reportlab(stripped[2:]), h1_style))
        elif stripped.startswith("## "):
            story.append(Paragraph(_sanitize_text_for_reportlab(stripped[3:]), h2_style))
        elif stripped.startswith("### "):
            story.append(Paragraph(_sanitize_text_for_reportlab(stripped[4:]), h2_style))
        else:
            story.append(Paragraph(_sanitize_text_for_reportlab(stripped), body_style))
            
    if code_lines:
        story.append(Preformatted("\n".join(code_lines), code_style))
        
    if not story:
        story.append(Paragraph("Empty Markdown Document", body_style))
        
    _build_pdf_from_elements(story, output_path)
    return output_path


def md_to_docx(input_path: str, output_path: str) -> str:
    """Convert Markdown to DOCX."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        md_text = f.read()
    doc = docx.Document()
    for line in md_text.splitlines():
        s = line.strip()
        if s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("- ") or s.startswith("* "):
            doc.add_paragraph(s[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)
    doc.save(output_path)
    return output_path


def md_to_txt(input_path: str, output_path: str) -> str:
    """Convert Markdown to plain text."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    # Strip basic markdown markers
    clean = re.sub(r'#+\s*', '', text)
    clean = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', clean)
    clean = re.sub(r'(\*\*|\*|__|_)', '', clean)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(clean)
    return output_path


def html_to_txt(input_path: str, output_path: str) -> str:
    """Strip HTML tags and write text."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    clean = re.sub(r'<style[^>]*>[\s\S]*?</style>', '', content, flags=re.IGNORECASE)
    clean = re.sub(r'<script[^>]*>[\s\S]*?</script>', '', clean, flags=re.IGNORECASE)
    clean = re.sub(r'<br\s*/?>', '\n', clean, flags=re.IGNORECASE)
    clean = re.sub(r'</p>', '\n\n', clean, flags=re.IGNORECASE)
    clean = re.sub(r'<[^>]+>', '', clean)
    clean = html.unescape(clean)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(clean.strip())
    return output_path


def html_to_docx(input_path: str, output_path: str) -> str:
    """Convert HTML document into DOCX."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    # Extract simple paragraphs and headers
    doc = docx.Document()
    tokens = re.findall(r'<(h[1-6]|p|li)[^>]*>(.*?)</\1>', content, flags=re.IGNORECASE | re.DOTALL)
    if tokens:
        for tag, text_inner in tokens:
            cleaned = html.unescape(re.sub(r'<[^>]+>', '', text_inner)).strip()
            if not cleaned:
                continue
            if tag.lower() == 'h1':
                doc.add_heading(cleaned, level=1)
            elif tag.lower() == 'h2':
                doc.add_heading(cleaned, level=2)
            elif tag.lower() == 'h3':
                doc.add_heading(cleaned, level=3)
            elif tag.lower() == 'li':
                doc.add_paragraph(cleaned, style='List Bullet')
            else:
                doc.add_paragraph(cleaned)
    else:
        # Fallback to plain text
        txt = re.sub(r'<[^>]+>', '', content)
        doc.add_paragraph(html.unescape(txt).strip())
    doc.save(output_path)
    return output_path


def html_to_pdf(input_path: str, output_path: str) -> str:
    """Convert HTML to PDF using reportlab."""
    with open(input_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()
    styles = getSampleStyleSheet()
    story = []
    tokens = re.findall(r'<(h[1-6]|p|li|pre)[^>]*>(.*?)</\1>', content, flags=re.IGNORECASE | re.DOTALL)
    if tokens:
        for tag, text_inner in tokens:
            cleaned = _sanitize_text_for_reportlab(re.sub(r'<[^>]+>', '', text_inner).strip())
            if not cleaned:
                continue
            if tag.lower() == 'h1':
                story.append(Paragraph(cleaned, styles['Heading1']))
            elif tag.lower() == 'h2':
                story.append(Paragraph(cleaned, styles['Heading2']))
            elif tag.lower() == 'h3':
                story.append(Paragraph(cleaned, styles['Heading3']))
            else:
                story.append(Paragraph(cleaned, styles['Normal']))
            story.append(Spacer(1, 6))
    else:
        plain = _sanitize_text_for_reportlab(re.sub(r'<[^>]+>', '', content).strip())
        story.append(Paragraph(plain or "Empty HTML", styles['Normal']))
    _build_pdf_from_elements(story, output_path)
    return output_path


def convert_document(input_path: str, source_ext: str, target_ext: str, output_dir: str) -> str:
    """Main document conversion dispatcher."""
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    out_file = os.path.join(output_dir, f"{base_name}.{target_ext}")

    pair = (source_ext, target_ext)

    # DOCX
    if pair == ("docx", "pdf"):
        return docx_to_pdf(input_path, out_file)
    elif pair == ("docx", "txt"):
        return docx_to_txt(input_path, out_file)
    elif pair == ("docx", "html"):
        return docx_to_html(input_path, out_file)

    # PDF
    elif pair == ("pdf", "docx"):
        return pdf_to_docx(input_path, out_file)
    elif pair == ("pdf", "txt"):
        return pdf_to_txt(input_path, out_file)
    elif pair == ("pdf", "html"):
        return pdf_to_html(input_path, out_file)
    elif source_ext == "pdf" and target_ext in ["png", "jpg"]:
        return pdf_to_images(input_path, output_dir, target_ext)

    # TXT
    elif pair == ("txt", "pdf"):
        return txt_to_pdf(input_path, out_file)
    elif pair == ("txt", "docx"):
        return txt_to_docx(input_path, out_file)
    elif pair == ("txt", "html"):
        return txt_to_html(input_path, out_file)

    # MD
    elif pair == ("md", "html"):
        return md_to_html(input_path, out_file)
    elif pair == ("md", "pdf"):
        return md_to_pdf(input_path, out_file)
    elif pair == ("md", "docx"):
        return md_to_docx(input_path, out_file)
    elif pair == ("md", "txt"):
        return md_to_txt(input_path, out_file)

    # HTML
    elif pair == ("html", "txt"):
        return html_to_txt(input_path, out_file)
    elif pair == ("html", "docx"):
        return html_to_docx(input_path, out_file)
    elif pair == ("html", "pdf"):
        return html_to_pdf(input_path, out_file)

    raise ValueError(f"Unsupported document conversion: {source_ext} -> {target_ext}")
