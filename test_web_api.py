"""
Integration tests for 56 File Converter Flask API & Endpoints.
"""

import io
import json
import unittest
from PIL import Image
import docx

from app import app


class TestWebApi(unittest.TestCase):

    def setUp(self):
        app.config['TESTING'] = True
        self.client = app.test_client()

    def test_index_page(self):
        """Verify homepage loads with title, theme switcher, share button, proper section order, and no FAQ."""
        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        html = res.data.decode("utf-8")
        self.assertIn("56 File Converter", html)
        self.assertIn("themeToggle", html)
        self.assertIn("shareBtn", html)
        self.assertIn("100% Free Forever", html)
        self.assertIn("No Daily Limits", html)
        self.assertIn("100% Secure", html)
        
        # Verify section order: How It Works -> Supported Formats -> Why Choose 56 File Converter
        pos_how = html.find('id="howItWorks"')
        pos_formats = html.find('id="formats"')
        pos_features = html.find('id="features"')
        self.assertTrue(pos_how != -1 and pos_formats != -1 and pos_features != -1)
        self.assertTrue(pos_how < pos_formats < pos_features)

        # Verify FAQ was removed
        self.assertNotIn("Frequently Asked Questions", html)

    def test_api_formats(self):
        """Verify formats API returns mapping."""
        res = self.client.get("/api/formats")
        self.assertEqual(res.status_code, 200)
        data = res.get_json()
        self.assertIn("conversion_map", data)
        self.assertIn("docx", data["conversion_map"])
        self.assertIn("pdf", data["conversion_map"]["docx"])
        self.assertIn("jpg", data["conversion_map"])
        self.assertIn("png", data["conversion_map"]["jpg"])

    def test_api_convert_image(self):
        """Verify JPG to PNG conversion via API."""
        img_io = io.BytesIO()
        img = Image.new("RGB", (64, 64), color=(255, 0, 0))
        img.save(img_io, format="JPEG")
        img_io.seek(0)

        data = {
            "file": (img_io, "test_pic.jpg"),
            "target_format": "png"
        }
        res = self.client.post("/api/convert", data=data, content_type="multipart/form-data")
        self.assertEqual(res.status_code, 200)
        json_resp = res.get_json()
        self.assertTrue(json_resp["success"])
        self.assertEqual(json_resp["target_format"], "png")
        self.assertIn("download_url", json_resp)

        # Download converted file
        dl_res = self.client.get(json_resp["download_url"])
        self.assertEqual(dl_res.status_code, 200)
        self.assertTrue(len(dl_res.data) > 0)

    def test_api_convert_docx_to_pdf(self):
        """Verify DOCX to PDF conversion via API."""
        doc = docx.Document()
        doc.add_heading("Web API Test", level=1)
        doc.add_paragraph("Testing docx to pdf upload through Flask API.")
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_io.seek(0)

        data = {
            "file": (docx_io, "doc_test.docx"),
            "target_format": "pdf"
        }
        res = self.client.post("/api/convert", data=data, content_type="multipart/form-data")
        self.assertEqual(res.status_code, 200)
        json_resp = res.get_json()
        self.assertTrue(json_resp["success"])
        self.assertEqual(json_resp["target_format"], "pdf")

        # Download PDF
        dl_res = self.client.get(json_resp["download_url"])
        self.assertEqual(dl_res.status_code, 200)
        self.assertTrue(len(dl_res.data) > 0)

    def test_api_batch_download_zip(self):
        """Verify ZIP generation for batch conversions."""
        # Create two conversions
        job_ids = []
        for i in range(2):
            img_io = io.BytesIO()
            img = Image.new("RGB", (32, 32), color=(0, 255, 0))
            img.save(img_io, format="PNG")
            img_io.seek(0)
            res = self.client.post(
                "/api/convert",
                data={"file": (img_io, f"img_{i}.png"), "target_format": "webp"},
                content_type="multipart/form-data"
            )
            job_ids.append(res.get_json()["job_id"])

        zip_res = self.client.post(
            "/api/download-all",
            data=json.dumps({"job_ids": job_ids}),
            content_type="application/json"
        )
        self.assertEqual(zip_res.status_code, 200)
        zip_json = zip_res.get_json()
        self.assertTrue(zip_json["success"])
        self.assertEqual(zip_json["file_count"], 2)

        # Download zip
        dl_zip = self.client.get(zip_json["zip_url"])
        self.assertEqual(dl_zip.status_code, 200)
        self.assertTrue(len(dl_zip.data) > 0)


if __name__ == "__main__":
    unittest.main()
